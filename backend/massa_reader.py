"""Leitura de planilhas de massa de teste em formato livre.

O QA monta a planilha como preferir, em BLOCOS: um título (ex.: "CEP NÃO CABEADO", "ALTO VALOR",
"CLIENTE MÓVEL APROVADO") e, embaixo, as massas — com ou sem linha de cabeçalho (CEP, NÚMERO, CPF…).
Sem cabeçalho, o campo sai do formato do valor (11 dígitos = CPF, 14 = CNPJ, 8 = CEP, até 5 = número).
Vários blocos podem ficar lado a lado na mesma aba.

Saída: {"abas": [{"nome", "blocos": [{"titulo", "campos", "linhas": [{campo: valor}], "celula"}]}]}
"""
import csv
import io
import re
import unicodedata

import openpyxl
from openpyxl.utils import get_column_letter

# Cabeçalho -> campo da massa (mesmos ids da aba Massa de Testes no navegador).
_CABECALHOS = {
    "cep": "cep",
    "numero": "numero", "num": "numero", "n": "numero", "no": "numero", "nro": "numero", "numero casa": "numero",
    "numero da casa": "numero", "numero residencia": "numero", "numero endereco": "numero",
    "cpf": "cpf", "documento": "cpf", "doc": "cpf", "cpf cnpj": "cpf",
    "cnpj": "cnpj",
    "contrato": "contrato", "codigo net": "contrato", "cod net": "contrato", "numero contrato": "contrato", "n contrato": "contrato",
    "nome": "nome", "nome completo": "nome", "razao social": "nome",
    "email": "email", "e mail": "email",
    "nascimento": "nascimento", "data nascimento": "nascimento", "dt nasc": "nascimento", "data de nascimento": "nascimento",
    "celular": "celular", "telefone": "celular", "fone": "celular", "msisdn": "celular", "ani": "celular",
}


def _norm(s):
    s = unicodedata.normalize("NFD", str(s or "")).encode("ascii", "ignore").decode("ascii").lower()
    return re.sub(r"[^a-z0-9]+", " ", s).strip()


def _texto(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def _campo_do_cabecalho(s):
    return _CABECALHOS.get(_norm(s))


def _e_dado(s):
    t = re.sub(r"[\s.\-/()]", "", s)
    return bool(t) and (t.isdigit() or "@" in s or re.fullmatch(r"\d{2}/\d{2}/\d{4}", s.strip()) is not None)


def _e_titulo(s):
    return bool(re.search(r"[A-Za-zÀ-ú]{3,}", s)) and _campo_do_cabecalho(s) is None and not _e_dado(s)


def _campo_pelo_valor(v, titulo):
    t = re.sub(r"\D", "", v)
    if "@" in v:
        return "email", v.strip()
    if re.fullmatch(r"\d{2}/\d{2}/\d{4}", v.strip()):
        return "nascimento", v.strip()
    if len(t) == 14:
        return "cnpj", t
    if len(t) == 8 and not re.search(r"cpf|cliente|movel", _norm(titulo)):
        return "cep", t
    if 9 <= len(t) <= 11:
        return "cpf", t.zfill(11)          # CPF sem o zero da frente (célula numérica)
    if len(t) <= 5:
        return "numero", t
    return "", t


def _normalizar_valor(campo, v):
    t = re.sub(r"\D", "", v)
    if campo == "cep" and t:
        return t.zfill(8)
    if campo == "cpf" and t and len(t) <= 11:
        return t.zfill(11)
    if campo in ("nome", "email", "nascimento"):
        return v.strip()
    return t or v.strip()


def _grade_xlsx(conteudo):
    wb = openpyxl.load_workbook(io.BytesIO(conteudo), data_only=True, read_only=True)
    for ws in wb.worksheets:
        linhas = [[_texto(c) for c in row] for row in ws.iter_rows(values_only=True)]
        yield ws.title, linhas


def _grade_csv(conteudo):
    texto = conteudo.decode("utf-8-sig", errors="replace")
    try:
        dialeto = csv.Sniffer().sniff(texto[:4000], delimiters=";,\t")
    except csv.Error:
        dialeto = csv.excel
        dialeto.delimiter = ";"
    yield "CSV", [[c.strip() for c in row] for row in csv.reader(io.StringIO(texto), dialeto)]


def _titulo_acima(cel, r, cols):
    """Texto-título mais próximo acima de um cabeçalho (até 3 linhas), nas colunas da tabela."""
    for rr in range(r - 1, max(-1, r - 4), -1):
        for cc in cols:
            s = cel(rr, cc)
            if s and _e_titulo(s):
                return s
    return ""


def _tabelas_da_grade(g, cel, alt, larg, usados):
    """Formato TABELA: uma linha de cabeçalho com pelo menos um campo de massa (CPF, CEP…) e, em cada linha
    de baixo, uma massa. As colunas de texto que não são campos (Perfil, Crédito, Grupo, Observação…) viram a
    DESCRIÇÃO daquela massa — é ela que liga a massa ao CT."""
    blocos = []
    for r in range(alt):
        cab = {c: cel(r, c) for c in range(larg) if cel(r, c) and (r, c) not in usados}
        campos = {c: _campo_do_cabecalho(t) for c, t in cab.items() if _campo_do_cabecalho(t)}
        if not campos:
            continue
        # Só é tabela se a linha de baixo tem dado em alguma coluna de campo.
        if not any(_e_dado(cel(r + 1, c)) for c in campos):
            continue
        # Colunas da tabela: o trecho contíguo do cabeçalho que contém os campos.
        cmin, cmax = min(campos), max(campos)
        while cmin - 1 >= 0 and cel(r, cmin - 1):
            cmin -= 1
        while cel(r, cmax + 1):
            cmax += 1
        cols = list(range(cmin, cmax + 1))
        atributos = {c: cab.get(c, "") for c in cols if c not in campos and cab.get(c)}
        titulo = _titulo_acima(cel, r, cols)
        linhas = []
        rr = r + 1
        while rr < alt and any(cel(rr, c) for c in cols):
            if any(_campo_do_cabecalho(cel(rr, c)) for c in campos) or (_e_titulo(cel(rr, cmin)) and not any(_e_dado(cel(rr, c)) for c in campos)):
                break
            linha, desc = {}, []
            for c in cols:
                v = cel(rr, c)
                if not v:
                    continue
                if c in campos and _e_dado(v):
                    linha[campos[c]] = _normalizar_valor(campos[c], v)
                elif c in atributos:
                    desc.append(f"{atributos[c]}: {v}")
                elif c in campos and campos[c] in ("nome",):
                    linha["nome"] = v
                usados.add((rr, c))
            if linha:
                if desc:
                    linha["_desc"] = "; ".join(desc)
                linhas.append(linha)
            rr += 1
        if linhas:
            for c in cols:
                usados.add((r, c))
            blocos.append({"titulo": titulo, "campos": sorted({k for ln in linhas for k in ln if not k.startswith("_")}),
                           "linhas": linhas, "celula": f"{get_column_letter(cmin + 1)}{r + 1}"})
    return blocos


def _soltos_da_grade(g, cel, alt, larg, usados):
    """O que sobrou (valores soltos, sem tabela nem bloco): reconhece pelo formato e usa como descrição o texto
    mais próximo à esquerda ou acima."""
    linhas = []
    for r in range(alt):
        for c in range(larg):
            v = cel(r, c)
            if not v or (r, c) in usados or not _e_dado(v):
                continue
            k, val = _campo_pelo_valor(v, "")
            if not k or k == "numero":
                continue
            desc = next((cel(r, cc) for cc in range(c - 1, max(-1, c - 4), -1) if cel(r, cc) and _e_titulo(cel(r, cc))), "") \
                or next((cel(rr, c) for rr in range(r - 1, max(-1, r - 6), -1) if cel(rr, c) and _e_titulo(cel(rr, c))), "")
            linha = {k: val}
            if desc:
                linha["_desc"] = desc
            linhas.append(linha)
            usados.add((r, c))
    return [{"titulo": "", "campos": sorted({k for ln in linhas for k in ln if not k.startswith("_")}), "linhas": linhas, "celula": "", "soltos": True}] if linhas else []


def _blocos_da_grade(g):
    alt = len(g)
    larg = max((len(r) for r in g), default=0)
    cel = lambda r, c: g[r][c] if 0 <= r < alt and 0 <= c < len(g[r]) else ""
    usados = set()
    blocos = _tabelas_da_grade(g, cel, alt, larg, usados)
    for r in range(alt):
        for c in range(larg):
            s = cel(r, c)
            # (células já lidas como tabela ficam de fora)
            if not s or (r, c) in usados or (r + 1, c) in usados or (r + 2, c) in usados or not _e_titulo(s):
                continue
            # Cabeçalho logo abaixo do título (opcional): define o campo de cada coluna.
            cab_r = r + 1
            campos_col = {}
            for cc in range(c, min(c + 6, larg)):
                campo = _campo_do_cabecalho(cel(cab_r, cc))
                if campo:
                    campos_col[cc] = campo
                elif cc > c and not cel(cab_r, cc):
                    break
            ini = cab_r + 1 if campos_col else r + 1
            if not campos_col:
                # Sem cabeçalho: colunas contíguas com dados na 1ª linha do bloco.
                cc = c
                while cc < larg and _e_dado(cel(ini, cc)):
                    campos_col[cc] = None
                    cc += 1
            if not campos_col or not any(_e_dado(cel(ini, cc)) for cc in campos_col):
                continue
            linhas = []
            rr = ini
            while rr < alt and any(_e_dado(cel(rr, cc)) for cc in campos_col):
                if _e_titulo(cel(rr, c)):
                    break
                linha = {}
                for cc, campo in campos_col.items():
                    v = cel(rr, cc)
                    if not v or not _e_dado(v):
                        continue
                    if campo:
                        linha[campo] = _normalizar_valor(campo, v)
                    else:
                        k, val = _campo_pelo_valor(v, s)
                        if k and k not in linha:
                            linha[k] = val
                    usados.add((rr, cc))
                if linha:
                    linhas.append(linha)
                rr += 1
            if linhas:
                usados.add((r, c))
                campos = sorted({k for ln in linhas for k in ln})
                blocos.append({"titulo": s, "campos": campos, "linhas": linhas, "celula": f"{get_column_letter(c + 1)}{r + 1}"})
    return blocos + _soltos_da_grade(g, cel, alt, larg, usados)


def ler_blocos(conteudo, nome_arquivo=""):
    nome = (nome_arquivo or "").lower()
    grades = _grade_csv(conteudo) if nome.endswith((".csv", ".txt")) else _grade_xlsx(conteudo)
    abas = []
    for aba, grade in grades:
        blocos = _blocos_da_grade(grade)
        if blocos:
            abas.append({"nome": aba, "blocos": blocos})
    return {"abas": abas, "total_blocos": sum(len(a["blocos"]) for a in abas)}
